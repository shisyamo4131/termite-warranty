from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OVERVIEW_PATH = ROOT / "white-ant-warranty-system-overview.docx"
RUNBOOK_PATH = ROOT / "white-ant-warranty-demo-runbook.docx"

BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9E1E8"
TEXT = RGBColor(45, 55, 65)
MUTED = RGBColor(90, 100, 110)
FONT = "Noto Sans CJK JP"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_styles(doc: Document, landscape: bool) -> None:
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.45)
        section.bottom_margin = Cm(1.35)
        section.left_margin = Cm(1.55)
        section.right_margin = Cm(1.55)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(24 if landscape else 22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    for name, size in (("Heading 1", 18), ("Heading 2", 14)):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("House Solution向け DEV環境デモ資料")
    set_run_font(run, 8, color=MUTED)


def add_title(doc: Document, title: str, subtitle: str, presenter: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_run_font(run, bold=True, color=RGBColor(0, 0, 0))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(subtitle)
    set_run_font(run, 13, color=MUTED)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Cm(4.0), Cm(12.0)]
    rows = [
        ("対象", "House Solution株式会社 清水社長"),
        ("説明担当", presenter),
        ("環境", "termite-warranty-dev DEV環境 架空デモデータ"),
    ]
    for row, values in zip(table.rows, rows):
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if idx == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, 9.5, bold=(idx == 0))
    set_table_borders(table)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 10.5, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, 10.5)


def add_image(doc: Document, filename: str, caption: str, width_inches: float, alt_text: str) -> None:
    path = SHOTS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(4)
    r = c.add_run(caption)
    set_run_font(r, 8.5, color=MUTED)


def add_two_images(doc: Document, left, right, width_inches: float = 5.0) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, item in zip(table.rows[0].cells, (left, right)):
        cell.width = Cm(13.0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=70, start=80, bottom=70, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(SHOTS / item[0]), width=Inches(width_inches))
        shape._inline.docPr.set("descr", item[2])
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(item[1])
        set_run_font(r, 8.2, color=MUTED)
    set_table_borders(table, color="FFFFFF")


def add_section_page(doc: Document, heading: str, summary: str, filename: str, caption: str) -> None:
    doc.add_page_break()
    doc.add_heading(heading, level=1)
    add_body(doc, summary)
    add_image(doc, filename, caption, 10.15, caption)


def build_overview() -> None:
    doc = Document()
    configure_styles(doc, landscape=True)
    add_title(
        doc,
        "白蟻保証業務管理システム 概要書",
        "現在の業務画面とデモで確認できる範囲",
        "株式会社AriLabo 中村氏",
    )
    add_body(
        doc,
        "本資料は、House Solutionの白蟻保証業務を管理するDEV環境プロトタイプの全体像を、実際の画面で説明するものです。案件、施主、物件、工務店、保証サービスを関連付け、保証期限と通知状態を一つのシステムで確認できます。",
    )
    add_body(doc, "画面内の名称、住所、電話番号、案件番号は説明用の架空データです。DEV環境は本番環境ではありません。")

    doc.add_heading("業務の流れ", level=1)
    flow = doc.add_table(rows=2, cols=5)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    steps = [
        ("1", "マスター登録"),
        ("2", "案件登録"),
        ("3", "保証管理"),
        ("4", "期限確認"),
        ("5", "工務店対応"),
    ]
    for i, (num, label) in enumerate(steps):
        top = flow.cell(0, i)
        bottom = flow.cell(1, i)
        set_cell_shading(top, BLUE)
        set_cell_shading(bottom, PALE_BLUE)
        for cell in (top, bottom):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=90, bottom=90)
        p = top.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(num)
        set_run_font(r, 11, bold=True, color=RGBColor(255, 255, 255))
        p = bottom.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, 10, bold=True)
    set_table_borders(flow)

    add_section_page(
        doc,
        "ダッシュボード",
        "ログイン直後に、表示対象の案件数、保証期限が30日以内の案件数、未通知の案件数を確認します。対応が必要な案件から、案件詳細へ直接移動できます。",
        "01-dashboard.png",
        "ダッシュボードでは最新20件を対象に現在の対応状況を確認します",
    )

    doc.add_page_break()
    doc.add_heading("案件の検索と確認", level=1)
    add_body(doc, "案件一覧には、案件番号、施主、物件住所、工務店、担当支店、状態が一行で表示されます。条件を指定すると、該当する案件を絞り込んで確認できます。")
    add_two_images(
        doc,
        ("03-case-list.png", "案件一覧", "案件一覧画面"),
        ("04-case-filter.png", "案件の絞り込み", "案件絞り込み画面"),
        width_inches=5.0,
    )

    add_section_page(
        doc,
        "案件登録",
        "案件登録では、物件を選択すると施主と工務店が初期選択されます。担当支店、申込日、引渡日、初回の保証サービスと保証開始日をまとめて登録できます。必要なマスターは入力途中でも追加できます。",
        "12-case-registration.png",
        "案件登録画面は初回保証までを一つの流れで入力します",
    )

    add_section_page(
        doc,
        "案件詳細と適用保証",
        "案件詳細では、案件に関連する物件、施主、工務店、担当支店、申込日、引渡日を確認できます。適用保証ごとに期間、開始日、満了日、通知状態、保証状態を保持します。保証を延長する場合は、既存の保証を残して新しい適用保証を追加します。",
        "05-case-detail.png",
        "案件詳細では案件情報と適用保証の履歴を確認します",
    )

    doc.add_page_break()
    doc.add_heading("業務マスター", level=1)
    add_body(doc, "工務店、施主、物件、保証サービスを個別に管理します。各一覧から詳細を開き、登録内容と関連情報を確認できます。無効化したマスターも履歴のために保持されます。")
    master_table = doc.add_table(rows=5, cols=3)
    master_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("マスター", "主な情報", "関連画面で確認できる内容")
    rows = [
        ("工務店", "住所、電話、担当者、備考", "担当物件と工務店アカウント"),
        ("施主", "住所、電話、FAX、備考", "所有物件"),
        ("物件", "住所、建築面積、施主、工務店", "案件登録で選択する物件情報"),
        ("保証サービス", "名称、種別、略称、標準保証期間", "有効な案件で使用中の対象物件"),
    ]
    for i, value in enumerate(headers):
        cell = master_table.cell(0, i)
        set_cell_shading(cell, BLUE)
        r = cell.paragraphs[0].add_run(value)
        set_run_font(r, 9.5, bold=True, color=RGBColor(255, 255, 255))
    for row_idx, values in enumerate(rows, 1):
        for col_idx, value in enumerate(values):
            cell = master_table.cell(row_idx, col_idx)
            if row_idx % 2 == 0:
                set_cell_shading(cell, "F6F9FC")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(value)
            set_run_font(r, 9.2, bold=(col_idx == 0))
    set_table_borders(master_table)
    add_two_images(
        doc,
        ("07-company-list.png", "工務店一覧", "工務店一覧画面"),
        ("06-warranty-service-list.png", "保証サービス一覧", "保証サービス一覧画面"),
        width_inches=4.85,
    )

    doc.add_page_break()
    doc.add_heading("マスター詳細と関連情報", level=1)
    add_body(doc, "工務店詳細では登録内容に加えて担当物件を確認できます。物件一覧では住所と建築面積を一覧で確認し、各物件の詳細へ移動できます。")
    add_two_images(
        doc,
        ("08-company-detail.png", "工務店詳細と担当物件", "工務店詳細画面"),
        ("11-property-list.png", "物件一覧", "物件一覧画面"),
        width_inches=4.9,
    )

    doc.add_page_break()
    doc.add_heading("アカウントと工務店連携", level=1)
    add_body(doc, "House Solution管理者は一般担当者と工務店の共通アカウントを管理できます。工務店からの仮申請や保証更改の回答は通知管理で確認します。工務店へのメールは、現在は送信待ちキューへの記録までで、外部のメール配送サービスには接続していません。")
    add_two_images(
        doc,
        ("09-notification-management.png", "工務店からの回答を扱う通知管理", "通知管理画面"),
        ("13-staff-account-management.png", "House Solution担当者のアカウント管理", "担当者アカウント管理画面"),
        width_inches=4.9,
    )

    doc.add_page_break()
    doc.add_heading("デモで確認する範囲", level=1)
    add_bullets(doc, [
        "DEV環境の架空データを使い、案件、適用保証、各マスター、通知管理、アカウント管理の画面を確認します。",
        "デモでは登録画面を開いて入力項目を説明しますが、データは登録・編集・無効化しません。",
        "本番環境の構築、外部メール配送、監視、バックアップと復旧は今回のデモ対象外です。",
        "画面上で確認された追加要望や仕様変更は、中村氏が別途取りまとめます。",
    ])
    add_body(doc, "資料基準日 2026年9月14日")
    doc.save(OVERVIEW_PATH)


def add_step(doc: Document, number: str, title: str, purpose: str, actions: list[str], talking: list[str], image: str | None = None) -> None:
    doc.add_page_break()
    doc.add_heading(f"{number} {title}", level=1)
    add_body(doc, f"目的  {purpose}")
    doc.add_heading("操作", level=2)
    add_numbered(doc, actions)
    doc.add_heading("説明する内容", level=2)
    add_bullets(doc, talking)
    if image:
        add_image(doc, image, title, 6.45, f"{title}の画面")


def build_runbook() -> None:
    doc = Document()
    configure_styles(doc, landscape=False)
    add_title(
        doc,
        "白蟻保証業務管理システム デモ実施手順書",
        "清水社長向け DEV環境デモの進行台本",
        "株式会社AriLabo 中村氏",
    )
    add_body(doc, "本手順書は、中村氏がDEV環境の架空データを使ってシステムの主要な業務画面を説明するための進行台本です。想定時間は20分から25分です。")

    info = doc.add_table(rows=5, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("URL", "https://termite-warranty-dev.web.app/"),
        ("ログインアカウント", "demo.admin@termite-warranty.com"),
        ("権限", "House Solution管理者"),
        ("使用データ", "DEV環境の架空デモデータ"),
        ("禁止操作", "登録確定、編集保存、無効化、アカウント発行"),
    ]
    for row, (label, value) in zip(info.rows, values):
        set_cell_shading(row.cells[0], PALE_BLUE)
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(label if idx == 0 else value)
            set_run_font(r, 9.5, bold=(idx == 0))
    set_table_borders(info)

    doc.add_heading("開始前の確認", level=1)
    add_bullets(doc, [
        "ログイン画面を開き、アカウントのメールアドレスとパスワードを入力しておきます。パスワードは本手順書に記録しません。",
        "DEV環境にデモ用データが表示されることを確認します。実在する顧客情報や物件情報は使用しません。",
        "デモ中は閲覧と画面遷移を中心にします。登録や無効化の確定ボタンは押しません。",
        "質問や要望はその場で判断せず、中村氏が持ち帰る項目として記録します。",
    ])

    add_step(
        doc, "1", "ログインとダッシュボード", "デモ環境と全体の入口を説明する", [
            "ログイン画面でメールアドレスが demo.admin@termite-warranty.com であることを確認します。",
            "パスワードが入力済みであることを確認し、ログインをクリックします。",
            "画面右上にデモ管理者と表示され、ダッシュボードが開くことを確認します。",
        ], [
            "この画面は本番環境ではなく、架空データを入れたDEV環境です。",
            "ダッシュボードでは、最新20件の案件、期限30日以内、未通知の件数を確認できます。",
            "対応が必要な案件から案件詳細へ移動できます。",
        ], "01-dashboard.png")

    add_step(
        doc, "2", "メニュー構成", "主要機能の配置を短時間で示す", [
            "画面左上のメニューボタンをクリックします。",
            "ダッシュボード、案件一覧、工務店、施主、保証サービス、物件、担当者アカウントを順に指し示します。",
            "工務店管理ポータルには通知管理とアカウント管理があることを説明します。",
            "案件一覧をクリックします。",
        ], [
            "案件を中心に、関連するマスターと工務店対応を同じシステムで扱います。",
            "今回のデモでは主要画面を確認し、データ変更は行いません。",
        ], "02-navigation.png")

    add_step(
        doc, "3", "案件一覧と絞り込み", "案件を探して詳細へ進む操作を示す", [
            "案件一覧の列を左から確認します。",
            "更新日時が新しい20件を表示していますという案内を指し示します。",
            "絞り込みをクリックします。",
            "案件番号、施主、物件、工務店、担当支店、保証サービス、保証種別、住所、通知状態、日付で絞り込めることを説明します。",
            "今回は条件を入力せず、キャンセルをクリックします。",
            "先頭に表示されている案件番号をクリックします。",
        ], [
            "一覧は一案件一行で、施主、物件住所、工務店、担当支店、状態を確認できます。",
            "未通知の適用保証がある案件には未通知と表示されます。",
            "条件を指定した場合は該当結果を20件ずつ確認できます。",
        ], "04-case-filter.png")

    add_step(
        doc, "4", "案件詳細と適用保証", "案件と複数保証の関係を説明する", [
            "案件番号、施主、工務店、案件状態を確認します。",
            "案件情報で物件、住所、担当支店、申込日、引渡日を確認します。",
            "適用保証のサービス、期間、開始日、満了日、通知、状態を確認します。",
            "編集と適用保証を追加のボタンは指し示すだけにし、クリックしません。",
            "一覧へ戻るをクリックします。",
        ], [
            "一つの案件に複数の適用保証を保持できます。",
            "保証期間と満了日は案件全体ではなく、適用保証ごとに管理します。",
            "保証を延長するときは既存保証を残し、新しい適用保証を追加します。",
        ], "05-case-detail.png")

    add_step(
        doc, "5", "案件登録画面", "新規案件の入力項目とマスター連携を説明する", [
            "案件一覧で新規登録をクリックします。",
            "物件、施主、工務店、担当支店、申込日、引渡日を上から確認します。",
            "初回保証の保証サービスと保証開始日を確認します。",
            "各選択欄の追加ボタンからマスターを追加できることを説明します。",
            "何も入力せず、キャンセルをクリックします。案件を登録はクリックしません。",
        ], [
            "物件を選ぶと、その物件に設定された施主と工務店が初期選択されます。",
            "必要なマスターがない場合も、案件入力を閉じずに追加できます。",
            "初回保証まで一つの登録操作で扱います。",
        ], "12-case-registration.png")

    add_step(
        doc, "6", "業務マスター", "案件登録を支える基本情報を示す", [
            "メニューから保証サービスを開きます。",
            "名称、種別、略称、標準保証期間を確認します。",
            "工務店を開き、住所、電話番号、担当者が一覧表示されることを確認します。",
            "先頭の工務店の詳細を開き、基本情報と担当物件を確認します。",
            "メニューから物件を開き、住所と建築面積を確認します。",
        ], [
            "工務店、施主、物件、保証サービスには個別の一覧と詳細があります。",
            "無効化は物理削除ではなく、過去の案件から参照できる状態で保持します。",
            "保証サービスの標準保証期間を変更しても、既存案件の適用保証は変わりません。",
        ], "06-warranty-service-list.png")

    add_step(
        doc, "7", "工務店とのやり取り", "工務店からの回答とアカウントの管理範囲を説明する", [
            "メニューの工務店管理ポータルから通知管理を開きます。",
            "仮申請 更改回答の一覧で種別、案件番号、工務店、物件と施主、状態を確認します。",
            "確認や更改依頼を作成のボタンはクリックしません。",
            "アカウント管理を開き、工務店ごとに共通アカウントを管理する画面であることを説明します。",
            "アカウントを発行と無効化はクリックしません。",
        ], [
            "工務店から提出された仮データを確認し、本登録または差し戻しを行う流れがあります。",
            "メール通知は現在、送信待ちキューへの記録までです。実際のメール配送には接続していません。",
            "工務店のパスワードはHouse Solution側で設定せず、工務店側が共通ログイン画面から設定します。",
        ], "09-notification-management.png")

    add_step(
        doc, "8", "担当者アカウント", "House Solution側の利用者管理を説明する", [
            "メニューから担当者アカウントを開きます。",
            "表示名、メールアドレス、権限、状態を確認します。",
            "新規発行、編集、無効化はクリックしません。",
            "メニューからダッシュボードへ戻ります。",
        ], [
            "House Solution管理者は一般担当者のアカウントを発行、編集、無効化できます。",
            "一般担当者のパスワードは本人が設定または再設定します。",
        ], "13-staff-account-management.png")

    doc.add_page_break()
    doc.add_heading("終了時の確認", level=1)
    add_bullets(doc, [
        "登録、編集、無効化、アカウント発行を確定していないことを確認します。",
        "ダッシュボードへ戻り、必要に応じてログアウトします。",
        "清水社長からの未決定事項への回答、仕様追加、変更要望は中村氏が別途整理します。",
        "画面表示が資料と異なる場合は、その場で推測せず、差分として持ち帰ります。",
    ])

    doc.add_heading("デモ中の注意", level=1)
    caution = doc.add_table(rows=5, cols=2)
    caution.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("データ変更", "確定操作を行わず、閲覧とキャンセルで進める"),
        ("個人情報", "実在する顧客、物件、連絡先を画面やメモに出さない"),
        ("メール", "送信済みとは説明せず、キュー記録までと説明する"),
        ("本番利用", "DEV環境のプロトタイプであり、本番稼働中とは説明しない"),
        ("未確認事項", "その場で断定せず、中村氏が持ち帰って整理する"),
    ]
    for row, (label, value) in zip(caution.rows, items):
        set_cell_shading(row.cells[0], PALE_BLUE)
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            r = cell.paragraphs[0].add_run(label if idx == 0 else value)
            set_run_font(r, 9.5, bold=(idx == 0))
    set_table_borders(caution)
    add_body(doc, "資料基準日 2026年9月14日")
    doc.save(RUNBOOK_PATH)


if __name__ == "__main__":
    build_overview()
    build_runbook()
    print(OVERVIEW_PATH)
    print(RUNBOOK_PATH)
