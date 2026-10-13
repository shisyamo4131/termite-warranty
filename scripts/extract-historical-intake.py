"""Read the reviewed XLSX into a private JSON input. Never prints property values."""
import argparse,datetime,hashlib,json
from pathlib import Path
import openpyxl

def extract(path, mapping):
    source_hash=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if source_hash!=mapping.get('sourceSha256'): raise ValueError('SOURCE_CHANGED')
    book=openpyxl.load_workbook(path,data_only=True)
    if book.sheetnames!=['HSC','SHS','UH']: raise ValueError('SHEET_NAMES')
    rows=[]
    def date(value):
        if not isinstance(value,(datetime.datetime,datetime.date)): raise ValueError('DATE_TYPE')
        return value.strftime('%Y-%m-%d')
    def text(value): return '' if value is None else str(value).strip()
    for group,limit in [('HSC',9),('SHS',18),('UH',8)]:
        sheet=book[group]
        for r in range(2,sheet.max_row+1):
            v=[sheet.cell(r,c).value for c in range(1,limit+1)]
            if not any(x is not None and x!='' for x in v): continue
            if v[0] is None: raise ValueError('UNEXPECTED_DATA_ROW')
            fields={}
            if group=='HSC':
                raw_branch=text(v[4]); number=text(v[0]); application=date(v[1]); handover=date(v[5])
                grade={'2等級':20,'3等級':30}.get(text(v[6]))
                fields=dict(name=text(v[2]),address=text(v[3]),buildingArea=v[8],foundationInsulation=True,insulationMaterials=text(v[7]),maxYears=grade)
                external=''
            elif group=='SHS':
                raw_branch=text(v[8]); number=text(v[14])+text(v[15]); application=date(v[1]); handover=date(v[12])
                grade={2:20,3:30}.get(v[4])
                if text(v[5]) not in ['あり','なし']: raise ValueError('INSULATION_VALUE')
                fields=dict(name=text(v[2]),address=text(v[3]),buildingArea=v[16],foundationInsulation=(v[5]=='あり' or r==62),insulationMaterials=text(v[6]),maxYears=grade,notes=text(v[17]))
                external=''
            else:
                raw_branch=text(v[5]); number=text(v[0]); application=date(v[1]); handover=date(v[6])
                if text(v[7]) not in ['あり','なし']: raise ValueError('INSULATION_VALUE')
                fields=dict(name=text(v[2]),address='',buildingArea=200,foundationInsulation=v[7]=='あり',insulationMaterials='',maxYears=10)
                if not isinstance(v[3],str): raise ValueError('EXTERNAL_KEY_TYPE')
                external=text(v[3])
            if fields['maxYears'] is None: raise ValueError('GRADE_VALUE')
            target=mapping['sourceBranches'].get(group,{}).get(raw_branch)
            if not target: raise ValueError('BRANCH_UNMAPPED')
            fields.update(applicationDate=application,handoverDate=handover)
            rows.append(dict(group=group,sourceRow=r,warrantyNumber=number,externalKey=external,companyId=target['companyId'],branchId=target['branchId'],fields=fields))
    counts={g:sum(r['group']==g for r in rows) for g in ['HSC','SHS','UH']}
    if counts!=mapping.get('expectedRows'): raise ValueError('ROW_COUNTS')
    return dict(schemaVersion=1,workbookId=mapping['workbookId'],sourceSha256=source_hash,transformVersion='historical-2026-10-v1',expectedRows=counts,decisions=mapping.get('decisions',{}),companies=mapping['companies'],branches=mapping['branches'],reservedNumbers=mapping.get('reservedNumbers',[]),rows=rows)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--xlsx',required=True);p.add_argument('--mapping',required=True);p.add_argument('--output',required=True);a=p.parse_args()
    try:
        output=extract(a.xlsx,json.loads(Path(a.mapping).read_text(encoding='utf-8')))
        # Exclusive creation prevents accidental replacement of a frozen input.
        with open(a.output,'x',encoding='utf-8') as f: json.dump(output,f,ensure_ascii=False,indent=2)
        print(json.dumps({'rows':len(output['rows']),'sourceSha256':output['sourceSha256']}))
    except Exception:
        raise SystemExit('EXTRACTION_FAILED: inspect the private source and mapping; no record values logged')
