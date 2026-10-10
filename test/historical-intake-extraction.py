import importlib.util,unittest,tempfile,json,hashlib,datetime,sys
from pathlib import Path
import openpyxl
sys.dont_write_bytecode=True
spec=importlib.util.spec_from_file_location('intake',Path(__file__).parents[1]/'scripts/extract-historical-intake.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
class ExtractionTests(unittest.TestCase):
    def test_source_changes_and_incomplete_counts_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'fictional.xlsx';b=openpyxl.Workbook();b.active.title='HSC';b.create_sheet('SHS');b.create_sheet('UH');b.save(path)
            mapping=dict(sourceSha256='a'*64,expectedRows=dict(HSC=99,SHS=66,UH=277))
            with self.assertRaisesRegex(ValueError,'SOURCE_CHANGED'):module.extract(path,mapping)
            mapping['sourceSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
            with self.assertRaisesRegex(ValueError,'ROW_COUNTS'):module.extract(path,mapping)
    def test_approved_shs_exception_dates_and_number_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'fictional.xlsx';b=openpyxl.Workbook();b.active.title='HSC';s=b.create_sheet('SHS');b.create_sheet('UH')
            values={1:datetime.datetime(2026,1,1),2:datetime.datetime(2026,6,1),3:'架空物件',4:'架空住所',5:3,6:'なし',7:'架空資材',9:'架空支店',13:datetime.datetime(2026,7,1),15:'SHS-',16:'26-001',17:200,18:'架空の送付メモ'}
            for c,v in values.items():s.cell(62,c,v)
            b.save(path)
            mapping=dict(sourceSha256=hashlib.sha256(path.read_bytes()).hexdigest(),expectedRows=dict(HSC=0,SHS=1,UH=0),sourceBranches=dict(SHS={'架空支店':dict(companyId='fictional',branchId='fictional')}),workbookId='fictional',companies=[],branches=[])
            output=module.extract(path,mapping);row=output['rows'][0]
            self.assertEqual(row['warrantyNumber'],'SHS-26-001');self.assertEqual(row['fields']['applicationDate'],'2026-06-01');self.assertTrue(row['fields']['foundationInsulation']);self.assertEqual(row['sourceRow'],62)
            self.assertEqual(row['fields']['notes'],'架空の送付メモ')
            s.move_range('A62:Q62',rows=1);b.save(path)
            with self.assertRaisesRegex(ValueError,'SOURCE_CHANGED'):module.extract(path,mapping)
if __name__=='__main__':unittest.main()
